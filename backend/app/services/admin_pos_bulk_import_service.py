import csv
import io
import unicodedata
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path

from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.models.branch import Branch
from app.models.pos_catalog import (
    PosBranchProduct,
    PosCategory,
    PosProduct,
)
from app.services.admin_pos_catalog_service import (
    AdminPosCatalogService,
)


class AdminPosBulkImportService:

    MAX_FILE_SIZE = 10 * 1024 * 1024
    MAX_ROWS = 2000

    HEADER_ALIASES = {
        "name": {
            "nombre",
            "name",
            "producto",
            "product",
        },
        "category": {
            "categoria",
            "category",
        },
        "sku": {
            "sku",
            "codigo",
            "codigo sku",
        },
        "base_price": {
            "precio base",
            "precio",
            "base price",
            "price",
        },
        "description": {
            "descripcion",
            "description",
        },
        "sort_order": {
            "orden",
            "order",
            "sort order",
        },
        "is_active": {
            "activo",
            "active",
            "is active",
            "estado",
        },
    }

    @staticmethod
    def _normalize(value) -> str:
        if value is None:
            return ""

        text = str(value).strip()

        text = unicodedata.normalize(
            "NFKD",
            text,
        )

        text = "".join(
            char
            for char in text
            if not unicodedata.combining(char)
        )

        text = (
            text
            .replace("_", " ")
            .replace("-", " ")
            .casefold()
        )

        return " ".join(text.split())

    @classmethod
    def _canonical_header(
        cls,
        value,
    ):
        normalized = cls._normalize(value)

        for canonical, aliases in cls.HEADER_ALIASES.items():
            if normalized in aliases:
                return canonical

        return None

    @staticmethod
    def _clean_text(value):
        if value is None:
            return None

        text = str(value).strip()

        return text or None

    @classmethod
    def _parse_bool(cls, value):
        if value is None or str(value).strip() == "":
            return True

        normalized = cls._normalize(value)

        if normalized in {
            "yes",
            "si",
            "true",
            "1",
            "activo",
            "active",
        }:
            return True

        if normalized in {
            "no",
            "false",
            "0",
            "inactivo",
            "inactive",
        }:
            return False

        raise ValueError(
            "Activo debe ser YES/NO, SI/NO o TRUE/FALSE."
        )

    @staticmethod
    def _parse_price(value):
        if value is None or str(value).strip() == "":
            raise ValueError(
                "Precio base requerido."
            )

        cleaned = (
            str(value)
            .replace("$", "")
            .replace(",", "")
            .strip()
        )

        try:
            amount = Decimal(cleaned)
        except InvalidOperation:
            raise ValueError(
                "Precio base inválido."
            )

        if amount < 0:
            raise ValueError(
                "Precio base no puede ser negativo."
            )

        return amount.quantize(
            Decimal("0.01"),
            rounding=ROUND_HALF_UP,
        )

    @staticmethod
    def _parse_order(value):
        if value is None or str(value).strip() == "":
            return 0

        try:
            return int(float(value))
        except (TypeError, ValueError):
            raise ValueError(
                "Orden debe ser un número entero."
            )

    @classmethod
    def _validate_file(
        cls,
        filename: str,
        content: bytes,
    ):
        if not content:
            raise HTTPException(
                422,
                "El archivo está vacío.",
            )

        if len(content) > cls.MAX_FILE_SIZE:
            raise HTTPException(
                413,
                "El archivo supera el máximo de 10 MB.",
            )

        suffix = Path(
            filename or ""
        ).suffix.lower()

        if suffix not in {
            ".xlsx",
            ".csv",
        }:
            raise HTTPException(
                422,
                "Formato no compatible. Usa XLSX o CSV.",
            )

        return suffix

    @classmethod
    def _read_xlsx(cls, content: bytes):
        try:
            workbook = load_workbook(
                io.BytesIO(content),
                read_only=True,
                data_only=True,
            )
        except Exception:
            raise HTTPException(
                422,
                "No se pudo leer el archivo Excel.",
            )

        sheet = workbook.active

        iterator = sheet.iter_rows(
            values_only=True
        )

        try:
            raw_headers = next(iterator)
        except StopIteration:
            raise HTTPException(
                422,
                "El archivo no contiene datos.",
            )

        headers = [
            cls._canonical_header(value)
            for value in raw_headers
        ]

        rows = []

        for excel_row_number, values in enumerate(
            iterator,
            start=2,
        ):
            if not any(
                value not in (None, "")
                for value in values
            ):
                continue

            row = {}

            for index, value in enumerate(values):
                if index >= len(headers):
                    continue

                header = headers[index]

                if header:
                    row[header] = value

            rows.append(
                (
                    excel_row_number,
                    row,
                )
            )

        return headers, rows

    @classmethod
    def _read_csv(cls, content: bytes):
        try:
            text = content.decode(
                "utf-8-sig"
            )
        except UnicodeDecodeError:
            try:
                text = content.decode(
                    "latin-1"
                )
            except UnicodeDecodeError:
                raise HTTPException(
                    422,
                    "No se pudo leer el archivo CSV.",
                )

        try:
            sample = text[:4096]

            dialect = csv.Sniffer().sniff(
                sample,
                delimiters=",;\t",
            )
        except csv.Error:
            dialect = csv.excel

        reader = csv.reader(
            io.StringIO(text),
            dialect,
        )

        try:
            raw_headers = next(reader)
        except StopIteration:
            raise HTTPException(
                422,
                "El archivo no contiene datos.",
            )

        headers = [
            cls._canonical_header(value)
            for value in raw_headers
        ]

        rows = []

        for csv_row_number, values in enumerate(
            reader,
            start=2,
        ):
            if not any(
                str(value).strip()
                for value in values
            ):
                continue

            row = {}

            for index, value in enumerate(values):
                if index >= len(headers):
                    continue

                header = headers[index]

                if header:
                    row[header] = value

            rows.append(
                (
                    csv_row_number,
                    row,
                )
            )

        return headers, rows

    @classmethod
    def _read_rows(
        cls,
        filename,
        content,
    ):
        suffix = cls._validate_file(
            filename,
            content,
        )

        if suffix == ".xlsx":
            headers, rows = cls._read_xlsx(
                content
            )
        else:
            headers, rows = cls._read_csv(
                content
            )

        required = {
            "name",
            "category",
            "base_price",
        }

        found = {
            header
            for header in headers
            if header
        }

        missing = required - found

        if missing:
            labels = {
                "name": "Nombre",
                "category": "Categoría",
                "base_price": "Precio base",
            }

            raise HTTPException(
                422,
                "Faltan columnas requeridas: "
                + ", ".join(
                    labels[item]
                    for item in sorted(missing)
                ),
            )

        if not rows:
            raise HTTPException(
                422,
                "No hay productos para importar.",
            )

        if len(rows) > cls.MAX_ROWS:
            raise HTTPException(
                422,
                f"Máximo {cls.MAX_ROWS} productos por archivo.",
            )

        return rows

    @classmethod
    def preview(
        cls,
        db: Session,
        filename: str,
        content: bytes,
    ) -> dict:
        raw_rows = cls._read_rows(
            filename,
            content,
        )

        categories = (
            db.execute(
                select(PosCategory)
            )
            .scalars()
            .all()
        )

        category_map = {
            cls._normalize(category.name):
                category
            for category in categories
        }

        existing_skus = {
            cls._normalize(sku)
            for sku in (
                db.execute(
                    select(PosProduct.sku).where(
                        PosProduct.sku.is_not(None)
                    )
                )
                .scalars()
                .all()
            )
            if sku
        }

        file_skus = set()
        result_rows = []

        for row_number, raw in raw_rows:
            errors = []

            name = cls._clean_text(
                raw.get("name")
            )

            category_name = cls._clean_text(
                raw.get("category")
            )

            sku = cls._clean_text(
                raw.get("sku")
            )

            description = cls._clean_text(
                raw.get("description")
            )

            category = None
            price = None
            sort_order = 0
            is_active = True

            if not name:
                errors.append(
                    "Nombre requerido."
                )
            elif len(name) > 160:
                errors.append(
                    "Nombre supera 160 caracteres."
                )

            if not category_name:
                errors.append(
                    "Categoría requerida."
                )
            else:
                category = category_map.get(
                    cls._normalize(
                        category_name
                    )
                )

                if category is None:
                    errors.append(
                        f'Categoría "{category_name}" no existe.'
                    )

            if sku:
                if len(sku) > 80:
                    errors.append(
                        "SKU supera 80 caracteres."
                    )

                normalized_sku = cls._normalize(
                    sku
                )

                if normalized_sku in existing_skus:
                    errors.append(
                        f'SKU "{sku}" ya existe.'
                    )

                if normalized_sku in file_skus:
                    errors.append(
                        f'SKU "{sku}" está duplicado en el archivo.'
                    )

                file_skus.add(
                    normalized_sku
                )

            try:
                price = cls._parse_price(
                    raw.get("base_price")
                )
            except ValueError as exc:
                errors.append(
                    str(exc)
                )

            try:
                sort_order = cls._parse_order(
                    raw.get("sort_order")
                )
            except ValueError as exc:
                errors.append(
                    str(exc)
                )

            try:
                is_active = cls._parse_bool(
                    raw.get("is_active")
                )
            except ValueError as exc:
                errors.append(
                    str(exc)
                )

            result_rows.append({
                "row_number": row_number,
                "valid": not errors,
                "errors": errors,

                "name": name,
                "category": category_name,
                "category_id": (
                    category.id
                    if category
                    else None
                ),

                "sku": sku,
                "base_price": (
                    str(price)
                    if price is not None
                    else None
                ),

                "description": description,
                "sort_order": sort_order,
                "is_active": is_active,
            })

        valid_count = sum(
            1
            for row in result_rows
            if row["valid"]
        )

        invalid_count = (
            len(result_rows)
            - valid_count
        )

        return {
            "filename": filename,
            "total_rows": len(result_rows),
            "valid_rows": valid_count,
            "invalid_rows": invalid_count,
            "can_import": (
                valid_count > 0
                and invalid_count == 0
            ),
            "rows": result_rows,
        }

    @classmethod
    def import_file(
        cls,
        db: Session,
        filename: str,
        content: bytes,
    ) -> dict:
        preview = cls.preview(
            db,
            filename,
            content,
        )

        if not preview["can_import"]:
            raise HTTPException(
                422,
                detail={
                    "message":
                        "Corrige los errores antes de importar.",
                    "preview": preview,
                },
            )

        branches = (
            db.execute(
                select(Branch).where(
                    Branch.is_active.is_(True)
                )
            )
            .scalars()
            .all()
        )

        try:
            for row in preview["rows"]:
                category = db.get(
                    PosCategory,
                    row["category_id"],
                )

                product = PosProduct(
                    category=category,
                    name=row["name"],
                    description=row["description"],
                    sku=row["sku"],
                    base_price=Decimal(
                        row["base_price"]
                    ),
                    image_url=None,
                    sort_order=row["sort_order"],
                    is_active=row["is_active"],
                )

                db.add(product)
                db.flush()

                for branch in branches:
                    db.add(
                        PosBranchProduct(
                            branch=branch,
                            product=product,
                            price=None,
                            is_available=True,
                            sort_order=0,
                        )
                    )

            db.commit()

        except IntegrityError:
            db.rollback()

            raise HTTPException(
                409,
                "Conflicto al importar. Revisa los SKU duplicados.",
            )

        except Exception:
            db.rollback()
            raise

        return {
            "created_count":
                preview["valid_rows"],

            "catalog":
                AdminPosCatalogService
                .get_catalog(db),
        }
